"""Paired ``v1`` / ``v2`` fixtures for ``xberg::diff::compare``.

Two scenarios:

- ``docx_memo_v1.docx`` vs ``docx_memo_v2.docx`` — same memo with one
  paragraph removed, one paragraph added, and one paragraph rewritten.
  Exercises ``ExtractionDiff.content_diff`` (DiffLine::Added / Removed).
- ``xlsx_budget_v1.xlsx`` vs ``xlsx_budget_v2.xlsx`` — a 3x3 budget table
  with one cell value changed in v2. Exercises ``ExtractionDiff.tables_changed``
  (the per-cell ``CellChange`` payload).

GT sidecars carry the same shape from ``gt_schema.diff_expectation``: the
relative paths to both halves of the pair, the substring assertions for
added/removed lines, and the expected ``CellChange`` entries (row/col/
from/to). Integration tests load BOTH halves, run extraction
independently, then call ``xberg::diff::compare`` and assert against
the GT.
"""

from __future__ import annotations

import io
from pathlib import Path

from docx import Document  # type: ignore[import-untyped, import-not-found, unused-ignore]
from openpyxl import Workbook  # type: ignore[import-untyped, import-not-found, unused-ignore]

from .gt_schema import diff_expectation, write_ground_truth

ZIP_MTIME = (2024, 1, 1, 0, 0, 0)

# DOCX content. Each entry is a single paragraph.
DOCX_V1 = [
    "Subject: Q2 planning meeting.",
    "Date: 2024-04-15.",
    "Attendees: Alice, Bob, Carol.",
    "Agenda item one: review last quarter's revenue.",
    "Agenda item two: discuss Q2 product launches.",
    "Action items will be circulated by Friday.",
]

DOCX_V2 = [
    "Subject: Q2 planning meeting.",
    "Date: 2024-04-15.",
    # "Attendees" line dropped in v2.
    "Agenda item one: review last quarter's revenue and margin.",  # rewritten
    "Agenda item two: discuss Q2 product launches.",
    "Agenda item three: hiring plan for engineering.",  # new
    "Action items will be circulated by Friday.",
]

# XLSX content. v2 changes B2 from 100 to 150.
XLSX_HEADER = ["Department", "Q1 Budget", "Q2 Budget"]
XLSX_V1_ROWS = [
    ["Engineering", 100, 120],
    ["Marketing", 50, 60],
    ["Operations", 80, 90],
]
XLSX_V2_ROWS = [
    ["Engineering", 150, 120],  # B2: 100 -> 150
    ["Marketing", 50, 60],
    ["Operations", 80, 90],
]


def _save_docx(paragraphs: list[str]) -> bytes:
    """Serialise a DOCX with one paragraph per entry."""
    doc = Document()
    for text in paragraphs:
        doc.add_paragraph(text)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _save_xlsx(header: list[str], rows: list[list[str | int]]) -> bytes:
    """Serialise a single-sheet workbook with ``header`` + ``rows``."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Budget"
    ws.append(header)
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _emit_docx_pair(output_dir: Path, repo_root: Path) -> list[Path]:
    v1_path = output_dir / "docx_memo_v1.docx"
    v2_path = output_dir / "docx_memo_v2.docx"
    sidecar_path = output_dir / "docx_memo_diff.gt.json"

    v1_path.write_bytes(_save_docx(DOCX_V1))
    v2_path.write_bytes(_save_docx(DOCX_V2))

    # Relative paths for the sidecar — both halves of the pair are needed
    # by the integration test.
    repo_root_resolved = repo_root.resolve()

    def _rel(path: Path) -> str:
        try:
            return str(path.resolve().relative_to(repo_root_resolved)).replace("\\", "/")
        except ValueError:
            return str(path.resolve()).replace("\\", "/")

    write_ground_truth(
        sidecar_path,
        v1_path,
        repo_root,
        document_format="docx",
        feature="diff",
        expectations=diff_expectation(
            before_path=_rel(v1_path),
            after_path=_rel(v2_path),
            content_changed=True,
            # Substrings that MUST appear in some DiffLine::Added entry.
            expected_added_lines=[
                "review last quarter's revenue and margin.",
                "Agenda item three: hiring plan for engineering.",
            ],
            # Substrings that MUST appear in some DiffLine::Removed entry.
            expected_removed_lines=[
                "Attendees: Alice, Bob, Carol.",
                "review last quarter's revenue.",
            ],
            notes=(
                "v2 drops the 'Attendees' line, rewrites agenda item one, and inserts "
                "agenda item three. Content paragraphs unchanged on either side stay "
                "in DiffLine::Context entries (not asserted)."
            ),
        ),
        generator="diff_pairs",
    )
    return [v1_path, v2_path, sidecar_path]


def _emit_xlsx_pair(output_dir: Path, repo_root: Path) -> list[Path]:
    v1_path = output_dir / "xlsx_budget_v1.xlsx"
    v2_path = output_dir / "xlsx_budget_v2.xlsx"
    sidecar_path = output_dir / "xlsx_budget_diff.gt.json"

    v1_path.write_bytes(_save_xlsx(XLSX_HEADER, XLSX_V1_ROWS))  # type: ignore[arg-type]
    v2_path.write_bytes(_save_xlsx(XLSX_HEADER, XLSX_V2_ROWS))  # type: ignore[arg-type]

    repo_root_resolved = repo_root.resolve()

    def _rel(path: Path) -> str:
        try:
            return str(path.resolve().relative_to(repo_root_resolved)).replace("\\", "/")
        except ValueError:
            return str(path.resolve()).replace("\\", "/")

    write_ground_truth(
        sidecar_path,
        v1_path,
        repo_root,
        document_format="xlsx",
        feature="diff",
        expectations=diff_expectation(
            before_path=_rel(v1_path),
            after_path=_rel(v2_path),
            content_changed=True,
            expected_added_lines=["150"],
            expected_removed_lines=["100"],
            table_cell_changes=[
                # Row 1 = Engineering row (header is row 0), col 1 = Q1 Budget.
                {"row": 1, "col": 1, "from": "100", "to": "150"},
            ],
            notes=(
                "Single cell change in B2 (Engineering / Q1 Budget): 100 -> 150. "
                "All other cells identical. ExtractionDiff.tables_changed should "
                "carry exactly one TableDiff with one CellChange entry."
            ),
        ),
        generator="diff_pairs",
    )
    return [v1_path, v2_path, sidecar_path]


def generate(output_root: Path, repo_root: Path) -> list[Path]:
    """Emit both diff pairs under ``output_root/diff/``."""
    output_dir = output_root / "diff"
    output_dir.mkdir(parents=True, exist_ok=True)
    written: list[Path] = []
    written.extend(_emit_docx_pair(output_dir, repo_root))
    written.extend(_emit_xlsx_pair(output_dir, repo_root))
    return written
